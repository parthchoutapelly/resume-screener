"""Document extraction Lambda (docs/02-ingestion-pipeline.md §6.2). Downloads
a JD/resume from S3, extracts plain text (no entity understanding at this
stage — Architecture.md §7), and synchronously invokes nlpProcessing.

Absolute prohibitions this file must never violate: no branching on filename,
no canned text, no fake OCR (every text path really runs Tesseract), no
fabricated confidence values (docs/02-ingestion-pipeline.md §1).
"""

from __future__ import annotations

import json
import os
import shutil
import uuid
import zipfile
from dataclasses import dataclass
from urllib.parse import unquote_plus

import boto3
import docx
import fitz  # PyMuPDF
import pytesseract
from botocore.config import Config
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader

from rs_common import clock, ids, log
from rs_common.ddb import conditional_update
from rs_common.errors import ErrorCode as C
from rs_common.errors import Stage as S
from rs_common.errors import TerminalError, TransientError, safe_message

MAX_FILE_BYTES = 10_485_760
MAX_PDF_PAGES = 10
OCR_DPI = 200
NATIVE_PAGE_MIN_CHARS = 40
MIN_USABLE_CHARS = 100
MAX_TEXT_CHARS = 100_000
RESUME_EXT = {"pdf", "docx", "png", "jpg", "jpeg", "tiff"}
JD_EXT = RESUME_EXT | {"txt"}  # txt only exists when the API wrote a pasted JD (D-45)

MAX_PAGE_DIMENSION_POINTS = 3000  # audit limit: max width/height in points (~41.6 in)
MAX_RASTER_PIXELS = 25_000_000  # audit limit: max rasterization pixel budget (25 MP)

Image.MAX_IMAGE_PIXELS = 50_000_000  # decompression-bomb guard (R-SEC-05)

s3 = boto3.client("s3")
lambda_client = boto3.client(
    "lambda",
    config=Config(read_timeout=40, connect_timeout=5, retries={"max_attempts": 0}),  # D-43
)
_ddb = boto3.resource("dynamodb")
JOBS = _ddb.Table(os.environ["JOBS_TABLE"])
CANDIDATES = _ddb.Table(os.environ["CANDIDATES_TABLE"])
FAILED = _ddb.Table(os.environ["FAILED_JOBS_TABLE"])
NLP_FUNCTION_NAME = os.environ["NLP_FUNCTION_NAME"]


@dataclass(frozen=True)
class Ref:
    doc_type: str
    job_id: str
    candidate_id: str | None
    key: str
    ext: str


def lambda_handler(event, context):
    for record in event["Records"]:  # BatchSize = 1
        receive_count = int(record.get("attributes", {}).get("ApproximateReceiveCount", "1"))
        body = json.loads(record["body"])
        if body.get("Event") == "s3:TestEvent":  # sent once when notifications are configured
            continue
        for s3rec in body.get("Records", []):
            process_object(
                s3rec["s3"]["bucket"]["name"],
                unquote_plus(s3rec["s3"]["object"]["key"]),
                receive_count,
            )


def parse_key(key: str) -> Ref | None:
    parts = key.split("/")
    if len(parts) == 3 and parts[0] == "jd-uploads":
        doc_type, job_id, cand, fname = "jd", parts[1], None, parts[2]
    elif len(parts) == 4 and parts[0] == "resume-uploads":
        doc_type, job_id, cand, fname = "resume", parts[1], parts[2], parts[3]
    else:
        return None
    ext = fname.rsplit(".", 1)[-1].lower() if "." in fname else ""
    return Ref(doc_type, job_id, cand, key, ext)


def process_object(bucket: str, key: str, receive_count: int) -> None:
    ref = parse_key(key)
    if ref is None:  # not ours — audit, never retry
        record_failure(
            None,
            S.ORPHAN_OBJECT,
            TerminalError(None, S.ORPHAN_OBJECT, "key matches no upload prefix"),
            receive_count,
            True,
        )
        return

    workdir = f"/tmp/{uuid.uuid4().hex}"  # R-SEC-05: never /tmp/{basename}
    os.makedirs(workdir)
    try:
        if not mark_ingest_started(ref):
            raise TerminalError(None, S.ORPHAN_OBJECT, "no placeholder item for this key")

        allowed = JD_EXT if ref.doc_type == "jd" else RESUME_EXT
        if ref.ext not in allowed:
            raise TerminalError(
                C.UNSUPPORTED_FORMAT, S.UNSUPPORTED_FORMAT, f"extension .{ref.ext} not allowed"
            )

        path = download(bucket, ref, workdir)
        text, meta = extract(path, ref.ext)

        if len("".join(text.split())) < MIN_USABLE_CHARS:
            raise TerminalError(
                C.UNREADABLE_DOCUMENT,
                S.OCR if meta.get("ocr_pages") else S.PDF_EXTRACT,
                f"only {len(text.strip())} usable chars",
            )

        invoke_nlp(ref, text, meta)
        log.info(
            "ingested",
            stage="extraction",
            job_id=ref.job_id,
            candidate_id=ref.candidate_id,
            doc_type=ref.doc_type,
            file_type=meta["file_type"],
            ocr_pages=meta.get("ocr_pages", 0),
        )
    except TerminalError as e:  # R-ERR-02: record, mark, ACK
        record_failure(ref, e.stage, e, receive_count, terminal=True)
        if e.code is not None:
            mark_error(ref, e.code)
    except TransientError as e:  # R-ERR-03: record, re-raise
        record_failure(ref, e.stage, e, receive_count, terminal=False)
        raise
    except Exception as e:  # unknown -> treat as transient
        stage = S.S3_DOWNLOAD if "botocore" in type(e).__module__ else S.PDF_EXTRACT
        record_failure(ref, stage, e, receive_count, terminal=False)
        raise
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


def download(bucket: str, ref: Ref, workdir: str) -> str:
    try:
        size = s3.head_object(Bucket=bucket, Key=ref.key)["ContentLength"]
    except Exception as e:
        raise TransientError(S.S3_DOWNLOAD, f"head_object failed: {type(e).__name__}") from e

    if size > MAX_FILE_BYTES:  # defence in depth; the POST policy already caps size
        raise TerminalError(C.FILE_TOO_LARGE, S.UNSUPPORTED_FORMAT, f"{size} bytes")

    path = os.path.join(workdir, f"source.{ref.ext}")
    try:
        s3.download_file(bucket, ref.key, path)
    except Exception as e:
        raise TransientError(S.S3_DOWNLOAD, f"download failed: {type(e).__name__}") from e
    return path


def extract(path: str, ext: str) -> tuple[str, dict]:
    if ext == "pdf":
        return extract_pdf(path)
    if ext == "docx":
        return extract_docx(path), {"file_type": "docx"}
    if ext == "txt":
        with open(path, encoding="utf-8", errors="replace") as f:
            return f.read(), {"file_type": "text"}
    return extract_image(path)


def extract_pdf(path: str) -> tuple[str, dict]:
    with open(path, "rb") as f:
        if b"%PDF-" not in f.read(1024):  # R-VAL-05 magic bytes
            raise TerminalError(C.UNSUPPORTED_FORMAT, S.UNSUPPORTED_FORMAT, "not a PDF")

    try:
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):  # owner-only passwords decrypt with ""
            raise TerminalError(C.UNREADABLE_DOCUMENT, S.PDF_EXTRACT, "password-protected PDF")
        n = len(reader.pages)
    except TerminalError:
        raise
    except Exception as e:
        raise TerminalError(
            C.UNREADABLE_DOCUMENT, S.PDF_EXTRACT, f"cannot open PDF: {type(e).__name__}"
        ) from e

    if n == 0:
        raise TerminalError(C.UNREADABLE_DOCUMENT, S.PDF_EXTRACT, "PDF has no pages")
    if n > MAX_PDF_PAGES:
        raise TerminalError(C.TOO_MANY_PAGES, S.PDF_EXTRACT, f"{n} pages")

    pages: list[str] = []
    for p in reader.pages:
        try:
            pages.append(p.extract_text() or "")
        except Exception:
            pages.append("")

    ocr_idx = [
        i for i, t in enumerate(pages) if len(t.strip()) < NATIVE_PAGE_MIN_CHARS
    ]  # per page (D-42, R-HON-04)
    if ocr_idx:
        try:
            with fitz.open(path) as doc:
                for i in ocr_idx:
                    page = doc[i]
                    rect = page.rect
                    calc_w = int(rect.width * OCR_DPI / 72)
                    calc_h = int(rect.height * OCR_DPI / 72)
                    if (
                        rect.width > MAX_PAGE_DIMENSION_POINTS
                        or rect.height > MAX_PAGE_DIMENSION_POINTS
                        or (calc_w * calc_h) > MAX_RASTER_PIXELS
                    ):
                        raise TerminalError(
                            C.UNREADABLE_DOCUMENT,
                            S.PDF_EXTRACT,
                            f"page {i+1} dimensions too large ({int(rect.width)}x{int(rect.height)})",
                        )
                    pix = page.get_pixmap(dpi=OCR_DPI, alpha=False)
                    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    pages[i] = pytesseract.image_to_string(img, lang="eng")  # real OCR (R-HON-03)
        except TerminalError:
            raise
        except Exception as e:
            raise TransientError(S.OCR, f"OCR failed: {type(e).__name__}") from e

    file_type = "pdf_native" if not ocr_idx else ("pdf_scanned" if len(ocr_idx) == n else "pdf_mixed")
    return "\n".join(pages), {"file_type": file_type, "page_count": n, "ocr_pages": len(ocr_idx)}


def extract_docx(path: str) -> str:
    if not zipfile.is_zipfile(path):
        raise TerminalError(C.UNSUPPORTED_FORMAT, S.UNSUPPORTED_FORMAT, "not a DOCX (not a zip)")
    with zipfile.ZipFile(path) as z:
        if "word/document.xml" not in z.namelist():
            raise TerminalError(C.UNSUPPORTED_FORMAT, S.UNSUPPORTED_FORMAT, "zip is not a Word document")
        if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:  # zip-bomb guard
            raise TerminalError(C.UNREADABLE_DOCUMENT, S.DOCX_EXTRACT, "uncompressed size too large")

    try:
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    except Exception as e:
        raise TerminalError(
            C.UNREADABLE_DOCUMENT, S.DOCX_EXTRACT, f"python-docx failed: {type(e).__name__}"
        ) from e


def extract_image(path: str) -> tuple[str, dict]:
    try:
        with Image.open(path) as im:
            im.verify()
        with Image.open(path) as im:
            if im.format not in {"PNG", "JPEG", "TIFF"}:
                raise TerminalError(C.UNSUPPORTED_FORMAT, S.UNSUPPORTED_FORMAT, f"image format {im.format}")
            frames = getattr(im, "n_frames", 1)  # multi-page TIFF
            if frames > MAX_PDF_PAGES:
                raise TerminalError(C.TOO_MANY_PAGES, S.OCR, f"{frames} frames")
            texts = []
            for i in range(frames):
                im.seek(i)
                texts.append(pytesseract.image_to_string(im.convert("RGB"), lang="eng"))
        return "\n".join(texts), {"file_type": "image", "page_count": frames, "ocr_pages": frames}
    except TerminalError:
        raise
    except (UnidentifiedImageError, Image.DecompressionBombError, SyntaxError, OSError) as e:
        code = C.UNSUPPORTED_FORMAT if isinstance(e, UnidentifiedImageError) else C.UNREADABLE_DOCUMENT
        raise TerminalError(code, S.OCR, f"image rejected: {type(e).__name__}") from e
    except Exception as e:
        raise TransientError(S.OCR, f"OCR failed: {type(e).__name__}") from e


def invoke_nlp(ref: Ref, text: str, meta: dict) -> None:
    payload = {
        "doc_type": ref.doc_type,
        "job_id": ref.job_id,
        "candidate_id": ref.candidate_id,
        "extracted_text": text[:MAX_TEXT_CHARS],
        "extraction_metadata": {
            "source_key": ref.key,
            "char_count": len(text),
            "text_truncated": len(text) > MAX_TEXT_CHARS,
            **meta,
        },
    }
    try:
        resp = lambda_client.invoke(
            FunctionName=NLP_FUNCTION_NAME,
            InvocationType="RequestResponse",
            Payload=json.dumps(payload).encode("utf-8"),
        )
    except Exception as e:
        raise TransientError(S.NLP, f"invoke failed: {type(e).__name__}") from e

    if resp.get("FunctionError"):
        err_type = json.loads(resp["Payload"].read() or b"{}").get("errorType", "Unknown")
        if err_type == "EmptyDocumentError":
            raise TerminalError(C.UNREADABLE_DOCUMENT, S.NLP, "no tokens after NLP")
        if err_type == "OrphanRecordError":
            raise TerminalError(None, S.ORPHAN_OBJECT, "item vanished before NLP write")
        raise TransientError(
            S.NLP, f"nlpProcessing raised {err_type}"
        )  # type only — message may contain text


def _item_key(ref: Ref):
    if ref.doc_type == "resume":
        return CANDIDATES, {"job_id": ref.job_id, "candidate_id": ref.candidate_id}
    return JOBS, {"job_id": ref.job_id}


def mark_ingest_started(ref: Ref) -> bool:
    table, key = _item_key(ref)
    now = clock.now_iso()
    return conditional_update(
        table,
        Key=key,
        UpdateExpression="SET ingest_started_at = if_not_exists(ingest_started_at, :t), updated_at = :t",
        ConditionExpression="attribute_exists(job_id)",
        ExpressionAttributeValues={":t": now},
    )


def mark_error(ref: Ref, code: C) -> None:
    table, key = _item_key(ref)
    ok = conditional_update(
        table,
        Key=key,
        UpdateExpression="SET parse_status = :e, error_code = :c, updated_at = :t",
        ConditionExpression="attribute_exists(job_id) AND parse_status <> :p",  # never regress `parsed` (Rules.md §5)
        ExpressionAttributeValues={":e": "error", ":c": str(code), ":p": "parsed", ":t": clock.now_iso()},
    )
    if not ok:
        log.info("status_not_regressed", stage="extraction", job_id=ref.job_id, candidate_id=ref.candidate_id)


def record_failure(ref: Ref | None, stage: S, exc: BaseException, retry_count: int, terminal: bool) -> None:
    item = {
        "job_id": ref.job_id if ref else "unknown",
        "failure_id": ids.failure_id(),
        "stage": str(stage),
        "error_type": type(exc).__name__,
        "error_message": safe_message(exc),
        "terminal": terminal,
        "retry_count": retry_count,
        "created_at": clock.now_iso(),
        "expires_at": clock.epoch_in_days(90),
    }
    if ref and ref.candidate_id:
        item["candidate_id"] = ref.candidate_id
    FAILED.put_item(Item=item)
    log.warning(
        "failure_recorded",
        stage=str(stage),
        terminal=terminal,
        job_id=item["job_id"],
        candidate_id=item.get("candidate_id"),
        error_type=item["error_type"],
    )
